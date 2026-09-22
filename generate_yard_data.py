from openpyxl import load_workbook
from collections import defaultdict

src = r'D:\STOCK\ALL CONTAINERS LOCATHIONS.xlsx'
output = r'D:\YMS\src\yardData.js'

wb = load_workbook(src, data_only=True, read_only=True)
ws = wb['المخزن الشامل (Master Data)']
rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[1] is not None]

valid_lines = {'A', 'B', 'C', 'D', 'E', 'F', 'G', 'W'}
blocks = defaultdict(list)

for row in rows:
    group = str(row[1]).strip().replace(' ', '') if row[1] is not None else ''
    if not group or 'إجمالي' in group or '-' not in group:
        continue

    line = group.split('-')[0].upper()
    if line not in valid_lines:
        continue

    block_num = group.split('-')[1]
    block_id = f'{line}{block_num}'
    exact = str(row[5]).strip() if row[5] is not None else ''
    if not exact:
        continue

    slot = exact.split('-')[-1]
    try:
        position = f'{block_id}-{int(slot):02d}'
    except ValueError:
        position = f'{block_id}-{slot.zfill(2)}'

    container_number = str(row[4]).strip() if row[4] is not None else ''
    if not container_number:
        container_number = str(row[2]).strip() + str(row[3]).strip()

    blocks[(line, block_id)].append({
        'position': position,
        'containerNumber': container_number,
        'size': '40ft',
        'status': 'Full',
        'truckPlate': 'Pending',
        'driverName': 'Pending',
        'timestamp': '2026-09-22T08:00',
    })

lines = []
for line in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'W']:
    line_blocks = []
    line_block_keys = [key for key in blocks if key[0] == line]
    line_block_keys.sort(key=lambda key: int(str(key[1])[1:]) if str(key[1])[1:].isdigit() else 999)

    for _, block_id in line_block_keys:
        items = sorted(blocks[(line, block_id)], key=lambda item: int(item['position'].split('-')[-1]))
        line_blocks.append({
            'id': block_id,
            'name': block_id,
            'capacity': 20,
            'containers': [
                {
                    'id': f"cont-{idx + 1}",
                    'position': item['position'],
                    'containerNumber': item['containerNumber'],
                    'size': item['size'],
                    'status': item['status'],
                    'truckPlate': item['truckPlate'],
                    'driverName': item['driverName'],
                    'timestamp': item['timestamp'],
                }
                for idx, item in enumerate(items)
            ],
        })

    lines.append({
        'id': line,
        'name': f'LINE {line}',
        'blocks': line_blocks,
    })

js_lines = []
for line in lines:
    block_strings = []
    for block in line['blocks']:
        container_strings = []
        for idx, item in enumerate(block['containers']):
            container_strings.append(
                "{ id: 'cont-" + str(idx + 1) + "', position: '" + item['position'] + "', containerNumber: '" + item['containerNumber'].replace("'", "\\'") + "', size: '" + item['size'] + "', status: '" + item['status'] + "', truckPlate: '" + item['truckPlate'].replace("'", "\\'") + "', driverName: '" + item['driverName'].replace("'", "\\'") + "', timestamp: '" + item['timestamp'] + "' }"
            )
        block_strings.append(
            "{ id: '" + block['id'] + "', name: '" + block['name'] + "', capacity: 20, containers: [" + ', '.join(container_strings) + "] }"
        )
    js_lines.append("{ id: '" + line['id'] + "', name: '" + line['name'] + "', blocks: [" + ', '.join(block_strings) + "] }")

content = "export const baseLines = [\n  " + ',\n  '.join(js_lines) + "\n]\n"
with open(output, 'w', encoding='utf-8') as f:
    f.write(content)

print(f'Generated {output}')
print(f'Lines: {len(lines)}')
print(f'Blocks: {sum(len(line["blocks"]) for line in lines)}')
print(f'Containers: {sum(len(block["containers"]) for line in lines for block in line["blocks"])}')
