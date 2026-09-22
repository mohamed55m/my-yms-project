from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = 'LINE A'

headers = ['Line', 'Block', 'Slot', 'Container No', 'Size', 'Truck Plate', 'Driver Name', 'Status', 'Gate-In Date']
ws.append(headers)

header_fill = PatternFill('solid', fgColor='D9EAF7')
for cell in ws[1]:
    cell.font = Font(bold=True, color='0B1F3A')
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

rows = [
    ('A', 'A1', 'A1-01', 'MSCU1234565', '40ft', 'B 1234 AB', 'Ali Rahman', 'Full', '2026-09-20 09:15'),
    ('A', 'A1', 'A1-02', 'MSCU2345676', '20ft', 'B 8856 DC', 'Nabil Saleh', 'Empty', '2026-09-20 12:10'),
    ('A', 'A2', 'A2-01', 'MSCU3456787', '40ft', 'B 2231 KK', 'Hamad Jaber', 'Full', '2026-09-20 16:35'),
    ('A', 'A3', 'A3-01', 'MSCU4567898', '20ft', 'B 7712 LL', 'Jamal Hadi', 'Empty', '2026-09-21 07:40'),
    ('A', 'A4', 'A4-01', 'MSCU5678909', '40ft', 'B 9910 MM', 'Fahad Nasser', 'Full', '2026-09-21 08:15'),
]
for row in rows:
    ws.append(row)

ws.freeze_panes = 'A2'
for column in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
    ws.column_dimensions[column].width = 18 if column in ['D', 'F', 'G', 'I'] else 12

wb.save('public/yard-template-A.xlsx')
print('Created: public/yard-template-A.xlsx')
